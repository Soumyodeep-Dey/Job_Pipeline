"""Read the supplied layouts without modifying or executing workbook content."""
from datetime import datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import or_, select
from sqlalchemy.exc import IntegrityError

from app.models import Company, DiscoveryConfig, Job
from app.schemas import JobCreate, name_key

TRACKER_HEADERS = [
    "Job ID", "Company", "Role", "Location", "Required Skills", "Match Score",
    "Missing Skills", "Source URL", "Date Found", "Status", "Resume Version", "Human Approval",
]
TRACKER_FIELDS = list(JobCreate.model_fields)
KEYWORD_SHEETS = ("Keywords for DATA", "Keywords for AI", "Keywords for Cyber")


def text(value):
    return str(value).strip() if value is not None else ""


def read_workbook(content: bytes):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Uncompressed workbook exceeds 100 MB")
        return load_workbook(BytesIO(content), data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("File must be a valid .xlsx workbook") from exc


def career_url(cell):
    value = cell.hyperlink.target if cell.hyperlink else text(cell.value)
    if value and not value.startswith(("https://", "http://")):
        raise ValueError(f"{cell.parent.title}!{cell.coordinate}: missing HTTP(S) career hyperlink")
    return value


def parse_companies(workbook):
    required = {"Company Data", *KEYWORD_SHEETS}
    if not required.issubset(workbook.sheetnames):
        raise ValueError(f"Missing sheets: {', '.join(sorted(required - set(workbook.sheetnames)))}")
    companies = {}

    def add(name, domain, url):
        key = name_key(name)
        item = companies.setdefault(key, dict(name=name, name_key=key, domains=[], career_pages=[], keyword_profiles=[]))
        if domain and domain not in item["domains"]:
            item["domains"].append(domain)
        if url and url not in item["career_pages"]:
            item["career_pages"].append(url)
        return item

    sheet = workbook["Company Data"]
    starts = [c.column for c in sheet[1] if c.value == "Company Name"]
    if starts != [1, 5, 9]:
        raise ValueError("Company Data: expected company blocks at A:C, E:G and I:K")
    source_entries = 0
    for row in range(2, sheet.max_row + 1):
        for col in starts:
            name = text(sheet.cell(row, col).value)
            if name:
                add(name, text(sheet.cell(row, col + 1).value), career_url(sheet.cell(row, col + 2)))
                source_entries += 1

    configs = []
    for sheet_name in KEYWORD_SHEETS:
        sheet = workbook[sheet_name]
        header_row = next((r for r in range(1, 11) if sheet.cell(r, 2).value == "Company"), None)
        bank_row = next((r for r in range(1, 11) if sheet.cell(r, 12).value == "Keyword / Phrase"), None)
        if header_row is None or bank_row is None:
            raise ValueError(f"{sheet_name}: company or keyword bank headers missing")
        if [sheet.cell(header_row, c).value for c in (6, 7)] != ["DATA", "EXCLUDE"]:
            raise ValueError(f"{sheet_name}: DATA / EXCLUDE columns missing")
        for row in range(header_row + 1, sheet.max_row + 1):
            name = text(sheet.cell(row, 2).value)
            if not name:
                continue
            domain = text(sheet.cell(row, 3).value)
            url = career_url(sheet.cell(row, 4))
            item = add(name, domain, url)
            item["keyword_profiles"].append({
                "sheet": sheet_name, "row": row, "domain": domain, "career_page": url,
                "pack": text(sheet.cell(row, 5).value),
                "keywords": [part.strip() for part in text(sheet.cell(row, 6).value).split(",") if part.strip()],
                "exclusions": [part.strip() for part in text(sheet.cell(row, 7).value).split(",") if part.strip()],
                "keyword_count": sheet.cell(row, 8).value,
                "guardrail": text(sheet.cell(row, 9).value),
            })
        bank = []
        for row in range(bank_row + 1, sheet.max_row + 1):
            values = [sheet.cell(row, c).value for c in range(11, 17)]
            if values[1] is not None:
                bank.append(dict(zip(("keyword_id", "phrase", "type", "area", "fit_tier", "recommended_use"), values)))
        # Preserve all rule, weight and source-reference rows verbatim, with coordinates.
        rules = [{"row": row, "values": [sheet.cell(row, c).value for c in range(18, 22)]}
                 for row in range(3, sheet.max_row + 1)
                 if any(sheet.cell(row, c).value is not None for c in range(18, 22))]
        notes = [[sheet.cell(row, c).value for c in range(1, sheet.max_column + 1)] for row in (1, 2)]
        configs.append(dict(sheet=sheet_name, keyword_bank=bank, rules=rules, notes=notes))
    return companies, configs, source_entries


def import_companies(session, workbook):
    companies, configs, entries = parse_companies(workbook)
    created = updated = 0
    for key, values in companies.items():
        company = session.scalar(select(Company).where(Company.name_key == key))
        if company is None:
            session.add(Company(**values))
            created += 1
        else:
            for field, value in values.items():
                setattr(company, field, value)
            updated += 1
    for values in configs:
        config = session.get(DiscoveryConfig, values["sheet"])
        if config is None:
            session.add(DiscoveryConfig(**values))
        else:
            for field, value in values.items():
                setattr(config, field, value)
    session.commit()
    return {"created": created, "updated": updated, "source_entries": entries,
            "unique_companies": len(companies),
            "keyword_banks": {c["sheet"]: len(c["keyword_bank"]) for c in configs}}


def build_job(session, payload):
    company = session.scalar(select(Company).where(Company.name_key == name_key(payload.company)))
    if company is None:
        raise ValueError(f"Unknown company '{payload.company}'; import companies first")
    values = payload.model_dump(exclude={"company"})
    return Job(company_id=company.id, **values)


def duplicate_job(session, payload):
    conditions = [Job.job_id == payload.job_id]
    if payload.source_url:
        conditions.append(Job.source_url == payload.source_url)
    return session.scalar(select(Job).where(or_(*conditions)))


def parse_approval(value):
    if value is None or text(value) == "":
        return False
    normalized = text(value).casefold()
    if normalized in ("true", "yes", "1"):
        return True
    if normalized in ("false", "no", "0"):
        return False
    raise ValueError("Human Approval must be yes/no, true/false, 1/0, or blank")


def import_jobs(session, workbook):
    sheets = [s for s in workbook if [text(c.value) for c in s[1]] == TRACKER_HEADERS]
    if len(sheets) != 1:
        raise ValueError("Expected exactly one sheet with the 12 tracker headers in their original order")
    sheet = sheets[0]
    result = {"created": 0, "duplicates": 0, "errors": [], "duplicate_rows": []}
    for row in sheet.iter_rows(min_row=2):
        if all(c.value is None or text(c.value) == "" for c in row):
            continue
        row_number = row[0].row
        try:
            if any(c.data_type == "f" for c in row):
                raise ValueError("Use values, not formulas, in tracker rows")
            values = {key: cell.value for key, cell in zip(TRACKER_FIELDS, row)
                      if cell.value is not None and text(cell.value) != ""}
            for field in ("job_id", "company", "role", "location", "required_skills", "missing_skills", "resume_version"):
                if field in values:
                    values[field] = text(values[field])
            if row[7].hyperlink:
                values["source_url"] = row[7].hyperlink.target
            if isinstance(values.get("date_found"), datetime):
                values["date_found"] = values["date_found"].date()
            values["human_approval"] = parse_approval(values.get("human_approval"))
            if not values.get("job_id") and not values.get("source_url"):
                raise ValueError("Excel rows need Job ID or Source URL for repeatable imports")
            payload = JobCreate(**values)
            if duplicate_job(session, payload):
                result["duplicates"] += 1
                result["duplicate_rows"].append(row_number)
                continue
            with session.begin_nested():
                session.add(build_job(session, payload))
                session.flush()
            result["created"] += 1
        except (ValueError, ValidationError) as exc:
            result["errors"].append({"row": row_number, "message": str(exc)})
        except IntegrityError:
            # A concurrent import may have inserted the same ID/URL after our check.
            if duplicate_job(session, payload):
                result["duplicates"] += 1
                result["duplicate_rows"].append(row_number)
            else:
                result["errors"].append({"row": row_number, "message": "Database constraint rejected this row"})
    session.commit()
    return result
